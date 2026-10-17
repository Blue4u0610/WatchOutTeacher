import base64
import io
import json
from pathlib import Path
import struct
import sys
import tempfile
import threading
import time
import unittest
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'bridge'))
from common import AppError
from host import Host, read_frame, encode_frame, extension_origin
from parser import parse
from providers import answer_with_fallback, validate, run, CLIProvider
from store import Store

GOOD={'answer':'B — Violet','explanation_en':'The course defines Violet as the readiness color.','sources':[{'source_id':'S1'}],'basis':'course_material','needs_more_context':False}
SOURCE={'S1':{'document_id':'d1','document':'handout.pdf','page_or_slide':2,'kind':'page'}}

class MockProvider:
    def __init__(self,result=None,code=None,delay=0): self.result=result;self.code=code;self.delay=delay;self.calls=0
    def answer(self,prompt,model,cancel,timeout):
        self.calls+=1
        if self.delay:
            cancel.wait(min(self.delay,timeout))
            if cancel.is_set(): raise AppError('cancelled','cancelled')
            if self.delay>=timeout: raise AppError('timeout','timeout')
        if self.code: raise AppError(self.code,self.code)
        return self.result

class ProviderTests(unittest.TestCase):
    def test_validate_citation_and_html_as_plain_string(self):
        r=validate(GOOD,SOURCE)
        self.assertEqual(r['sources'][0]['page_or_slide'],2)
        with self.assertRaises(AppError): validate(GOOD,{})
        with self.assertRaises(AppError): validate({**GOOD,'needs_more_context':'false'},SOURCE)
        with self.assertRaises(AppError): validate({**GOOD,'answer':''},SOURCE)
    def test_fallback_classified_failures(self):
        for code in ['not_installed','auth','rate_limit','timeout','format','provider']:
            p={'codex':MockProvider(code=code),'claude':MockProvider(result=GOOD)}
            r=answer_with_fallback(p,'codex',True,'',{},threading.Event(),3,SOURCE)
            self.assertEqual(r['diagnostics']['provider'],'claude')
            self.assertIn(code,r['diagnostics']['fallback_reason'])
    def test_no_fallback_on_cancellation(self):
        p={'codex':MockProvider(code='cancelled'),'claude':MockProvider(result=GOOD)}
        with self.assertRaises(AppError): answer_with_fallback(p,'codex',True,'',{},threading.Event(),3,SOURCE)
        self.assertEqual(p['claude'].calls,0)
    def test_normal_insufficient_not_failure(self):
        p={'codex':MockProvider(result={**GOOD,'basis':'insufficient','sources':[],'needs_more_context':True}),'claude':MockProvider(code='not_installed')}
        r=answer_with_fallback(p,'codex',True,'',{},threading.Event(),3,{})
        self.assertEqual(r['basis'],'insufficient'); self.assertEqual(p['claude'].calls,0)
    def test_total_timeout_budget(self):
        p={n:MockProvider(delay=3) for n in ['codex','claude']}
        start=time.monotonic()
        with self.assertRaises(AppError): answer_with_fallback(p,'codex',True,'',{},threading.Event(),.15,{})
        self.assertLess(time.monotonic()-start,.35)
    def test_both_missing_and_fallback_disabled(self):
        p={n:MockProvider(code='not_installed') for n in ['codex','claude']}
        with self.assertRaisesRegex(AppError,'codex.*claude'): answer_with_fallback(p,'codex',True,'',{},threading.Event(),3,{})
        p['claude'].calls=0
        with self.assertRaises(AppError): answer_with_fallback(p,'codex',False,'',{},threading.Event(),3,{})
        self.assertEqual(p['claude'].calls,0)
    def test_real_subprocess_cancel_and_timeout(self):
        with tempfile.TemporaryDirectory() as d:
            start=time.monotonic()
            with self.assertRaises(AppError) as e: run([sys.executable,'-c','import time; time.sleep(10)'],d,timeout=.15)
            self.assertEqual(e.exception.code,'timeout'); self.assertLess(time.monotonic()-start,1)
            event=threading.Event(); timer=threading.Timer(.1,event.set); timer.start()
            with self.assertRaises(AppError) as e: run([sys.executable,'-c','import time; time.sleep(10)'],d,cancel=event,timeout=5)
            self.assertEqual(e.exception.code,'cancelled'); timer.join()
    def test_format_failure_falls_back(self):
        p={'codex':MockProvider(result={'progress':'loading'}),'claude':MockProvider(result=GOOD)}
        self.assertEqual(answer_with_fallback(p,'codex',True,'',{},threading.Event(),2,SOURCE)['diagnostics']['provider'],'claude')

class ProtocolTests(unittest.TestCase):
    def test_fragmented_and_multiple_frames(self):
        class Fragmented(io.BytesIO):
            def read(self,n=-1): return super().read(min(n,2))
        stream=Fragmented(encode_frame({'id':'1','题':'你好'})+encode_frame({'id':'2'}))
        self.assertEqual(read_frame(stream)['题'],'你好'); self.assertEqual(read_frame(stream)['id'],'2'); self.assertIsNone(read_frame(stream))
    def test_invalid_frames(self):
        for data in [b'\x01',struct.pack('=I',9999999),struct.pack('=I',5)+b'{}',struct.pack('=I',2)+b'xx']:
            with self.assertRaises(AppError): read_frame(io.BytesIO(data))
    def test_origin_pinned(self): self.assertRegex(extension_origin(),r'^chrome-extension://[a-p]{32}/$')

class StorageTests(unittest.TestCase):
    def setUp(self): self.tmp=tempfile.TemporaryDirectory(); self.store=Store(Path(self.tmp.name))
    def tearDown(self): self.tmp.cleanup()
    def test_course_isolation_delete_index_and_persistence(self):
        a=self.store.create('A')['settings']['active_course']
        s=self.store.add(a,'first.pdf',{'units':[{'number':12,'kind':'page','text':'Lumen Violet ready'}],'warnings':[]})
        doc=s['documents'][0]['id']
        b=self.store.create('B')['settings']['active_course']
        self.store.add(b,'second.pptx',{'units':[{'number':3,'kind':'slide','text':'Lumen Amber ready'}],'warnings':[]})
        r,_,_=self.store.retrieve(a,'Lumen'); self.assertEqual(r[0]['document_id'],doc); self.assertEqual(r[0]['page_or_slide'],12)
        r,_,_=self.store.retrieve(b,'Lumen'); self.assertEqual(r[0]['text'],'Lumen Amber ready')
        self.store.remove(doc); self.assertEqual(self.store.retrieve(a,'Lumen')[0],[])
        self.store.rename(b,'Renamed'); self.assertEqual(Store(Path(self.tmp.name)).state()['courses'][1]['name'],'Renamed')
        self.store.delete(b); self.assertIsNone(self.store.settings()['active_course'])
    def test_retrieval_is_bounded_and_diagnostic(self):
        a=self.store.create('A')['settings']['active_course']
        self.store.add(a,'large.txt',{'units':[{'number':1,'kind':'text','text':'Lumen ready '*5000}],'warnings':[]})
        r,d,w=self.store.retrieve(a,'Lumen'); self.assertEqual(len(r),12); self.assertGreater(d['total'],12); self.assertTrue(w)
    def test_effort_setting_persists(self):
        self.assertEqual(self.store.settings()['codex_effort'],'')
        self.store.save_settings({'codex_effort':'high'})
        self.assertEqual(Store(self.store.directory).settings()['codex_effort'],'high')

    def test_settings_validation(self):
        for payload in [{'primary':'x'},{'timeout':1},{'codex_model':'hello; evil'},{'executable':'/bin/sh'},{'fallback':'yes'},{'codex_effort':'ultra'},{'codex_effort':['high']},{'codex_effort':'high; evil'}]:
            with self.assertRaises(AppError): self.store.save_settings(payload)

class ImportTests(unittest.TestCase):
    def setUp(self): self.tmp=tempfile.TemporaryDirectory(); self.path=Path(self.tmp.name)
    def tearDown(self): self.tmp.cleanup()
    def test_text_and_invalid_formats(self):
        p=self.path/'test.txt'; p.write_text('Real body.'); self.assertEqual(parse(p,'.txt')['units'][0]['text'],'Real body.')
        p.write_text('');
        with self.assertRaises(AppError): parse(p,'.txt')
        with self.assertRaises(AppError): parse(p,'.ppt')
        p.write_bytes(b'corrupt')
        with self.assertRaises(AppError): parse(p,'.pdf')
        with self.assertRaises(AppError): parse(p,'.pptx')
    def test_scanned_and_encrypted_pdf(self):
        from pypdf import PdfWriter
        p=self.path/'scan.pdf'; w=PdfWriter(); w.add_blank_page(width=200,height=200); w.write(str(p))
        with self.assertRaises(AppError) as e: parse(p,'.pdf')
        self.assertEqual(e.exception.code,'ocr_required')
        w.encrypt('secret'); w.write(str(p))
        with self.assertRaises(AppError) as e: parse(p,'.pdf')
        self.assertEqual(e.exception.code,'encrypted')
    def test_pdf_page_numbers_and_pptx_slides(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
        w=PdfWriter(); w.add_blank_page(width=200,height=200); page=w.add_blank_page(width=200,height=200)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
        stream=DecodedStreamObject(); stream.set_data(b'BT /F1 12 Tf 10 150 Td (Lumen Violet ready) Tj ET');page[NameObject('/Contents')]=w._add_object(stream)
        p=self.path/'actual.pdf'; w.write(str(p)); parsed=parse(p,'.pdf'); self.assertIn('Lumen',parsed['units'][1]['text']); self.assertEqual(parsed['units'][1]['number'],2)
        from pptx import Presentation
        from pptx.util import Inches
        deck=Presentation();deck.slides.add_slide(deck.slide_layouts[6]);slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(Inches(1),Inches(1),Inches(5),Inches(1)).text='Lumen readiness is Violet'
        p=self.path/'actual.pptx'; deck.save(str(p)); parsed=parse(p,'.pptx'); self.assertEqual(parsed['units'][1]['number'],2); self.assertIn('Violet',parsed['units'][1]['text'])
    def test_chunk_upload_and_failed_status(self):
        h=Host(self.path)
        try:
            course=h.dispatch('course_create',{'name':'Test'})['settings']['active_course']
            data=b'Lumen protocol readiness color is Violet.'
            token=h.dispatch('upload_begin',{'name':'lesson.md','size':len(data),'course':course})['token']
            with self.assertRaises(AppError): h.dispatch('upload_chunk',{'token':token,'offset':3,'data':base64.b64encode(data).decode()})
            h.dispatch('upload_chunk',{'token':token,'offset':0,'data':base64.b64encode(data).decode()})
            state=h.dispatch('upload_finish',{'token':token}); self.assertEqual(state['documents'][0]['status'],'ready')
            token=h.dispatch('upload_begin',{'name':'broken.pdf','size':3,'course':course})['token']
            h.dispatch('upload_chunk',{'token':token,'offset':0,'data':base64.b64encode(b'bad').decode()})
            with self.assertRaises(AppError):h.dispatch('upload_finish',{'token':token})
            self.assertEqual(h.store.state()['documents'][1]['status'],'error');self.assertEqual(h.uploads,{})
        finally:h.cleanup()

class AdapterContractTests(unittest.TestCase):
    """Fixture outputs; these are NOT live Claude integration tests."""
    def test_model_catalog_filters_hidden_invalid_and_duplicate_rows(self):
        from unittest.mock import patch
        rows=[{'slug':'public-model','display_name':'Public','visibility':'list','supported_reasoning_levels':[{'effort':'low'},{'effort':'high'},{'effort':'ultra'},{'effort':'invalid'},None],'default_reasoning_level':'low'},
              {'slug':'hidden-model','visibility':'hide'},
              {'slug':'public-model','visibility':'list'},
              {'slug':'--unsafe','visibility':'list'},None,
              {'slug':'second','visibility':'list'}]
        with patch('providers.locate',return_value='/trusted/codex'),patch('providers.run',return_value=json.dumps({'models':rows})) as child:
            result=CLIProvider('codex').models()
            self.assertEqual(result['models'],[{'id':'public-model','name':'Public','efforts':['low','high'],'default_effort':'low'},{'id':'second','name':'second','efforts':[],'default_effort':None}])
            self.assertEqual(child.call_args.args[0],['/trusted/codex','debug','models'])
    def test_model_catalog_failure_and_missing_cli(self):
        from unittest.mock import patch
        with patch('providers.locate',return_value='/trusted/codex'),patch('providers.run',return_value='{"models":null}'):
            with self.assertRaises(AppError):CLIProvider('codex').models()
        with patch('providers.locate',return_value=None),patch('providers.run') as child:
            self.assertEqual(CLIProvider('codex').models()['models'],[])
            child.assert_not_called()

    def test_codex_effort_explicit_argument_and_auto_omission(self):
        from unittest.mock import patch
        p=CLIProvider('codex')
        for effort in ['', 'low','medium','high','xhigh','max']:
            def child(args,cwd,*rest):
                overrides=[args[i+1] for i,a in enumerate(args[:-1]) if a=='-c' and args[i+1].startswith('model_reasoning_effort=')]
                self.assertEqual(overrides,['model_reasoning_effort='+json.dumps(effort)] if effort else [])
                Path(args[args.index('--output-last-message')+1]).write_text(json.dumps(GOOD))
                return ''
            with patch.object(p,'detect',return_value={'available':True,'path':'/trusted/codex'}),patch('providers.run',side_effect=child):
                self.assertEqual(p.answer('question','model',threading.Event(),3,effort=effort),GOOD)
        with patch.object(p,'detect') as detect:
            with self.assertRaises(AppError):p.answer('question','model',threading.Event(),3,effort='ultra')
            detect.assert_not_called()

    def test_effort_forwarding_and_fallback_isolation(self):
        from unittest.mock import Mock
        codex=Mock();claude=Mock()
        codex.answer.return_value=GOOD
        result=answer_with_fallback({'codex':codex,'claude':claude},'codex',True,'q',{'codex':'m'},threading.Event(),3,SOURCE,efforts={'codex':'high'})
        self.assertEqual(codex.answer.call_args.kwargs,{'effort':'high'})
        self.assertEqual(result['diagnostics']['requested_effort'],'high')
        codex.answer.side_effect=AppError('provider','failed');claude.answer.return_value=GOOD
        result=answer_with_fallback({'codex':codex,'claude':claude},'codex',True,'q',{},threading.Event(),3,SOURCE,efforts={'codex':'high'})
        self.assertEqual(claude.answer.call_args.kwargs,{})
        self.assertEqual(result['diagnostics']['requested_effort'],'auto')

    def test_claude_envelope_and_restriction_arguments(self):
        from unittest.mock import patch
        p=CLIProvider('claude')
        with patch.object(p,'detect',return_value={'available':True,'path':'/trusted/claude'}), patch('providers.run',return_value=json.dumps({'type':'result','is_error':False,'structured_output':GOOD})) as child:
            result=p.answer('UNTRUSTED question','sonnet',threading.Event(),3)
            self.assertEqual(result,GOOD)
            args=child.call_args.args[0]
            self.assertEqual(args[args.index('--tools')+1],'')
            for flag in ['--safe-mode','--no-session-persistence','--strict-mcp-config','--no-chrome']:self.assertIn(flag,args)
            self.assertNotIn('UNTRUSTED question',args)
            self.assertEqual(child.call_args.args[2],'UNTRUSTED question')
    def test_claude_progress_is_not_an_answer(self):
        from unittest.mock import patch
        p=CLIProvider('claude')
        with patch.object(p,'detect',return_value={'available':True,'path':'/trusted/claude'}), patch('providers.run',return_value=json.dumps({'type':'assistant','text':'a progress line'})):
            with self.assertRaises(AppError):p.answer('question','',threading.Event(),3)
    def test_codex_only_final_file_and_tools_disabled(self):
        from unittest.mock import patch
        p=CLIProvider('codex')
        def child(args,cwd,*rest):
            Path(args[args.index('--output-last-message')+1]).write_text(json.dumps(GOOD))
            self.assertIn('--ignore-user-config',args);self.assertIn('--ephemeral',args);self.assertIn('read-only',args)
            self.assertIn('shell_tool',args);self.assertIn('hooks',args)
            return '{"type":"item.completed","item":{"type":"agent_message","text":"NOT THE ANSWER"}}'
        with patch.object(p,'detect',return_value={'available':True,'path':'/trusted/codex'}),patch('providers.run',side_effect=child):
            self.assertEqual(p.answer('question','',threading.Event(),3),GOOD)
    def test_subprocess_stdout_limit_after_fast_exit(self):
        with tempfile.TemporaryDirectory() as d:
            with self.assertRaises(AppError) as e:run([sys.executable,'-c','print("x"*20000)'],d,output_limit=1000)
            self.assertEqual(e.exception.code,'output_limit')
    def test_capability_lock_wait_obeys_budget_and_cancel(self):
        from providers import cancellable_lock
        lock=threading.Lock();lock.acquire()
        try:
            with self.assertRaises(AppError) as e:
                with cancellable_lock(lock,threading.Event(),time.monotonic()+.05):pass
            self.assertEqual(e.exception.code,'timeout')
            event=threading.Event();event.set()
            with self.assertRaises(AppError) as e:
                with cancellable_lock(lock,event,time.monotonic()+3):pass
            self.assertEqual(e.exception.code,'cancelled')
        finally:lock.release()

if __name__=='__main__': unittest.main()
