"""Generate original, tiny PDF/PPTX fixtures with known page/slide numbers."""
import sys
from pathlib import Path
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
from pptx import Presentation
from pptx.util import Inches
out=Path(sys.argv[1]);out.mkdir(parents=True,exist_ok=True)
lines=['Lumen Protocol Handout','Lumen readiness color is Violet. Amber means waiting. Green means shutdown.']
w=PdfWriter()
for line in lines:
    page=w.add_blank_page(width=620,height=300)
    font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
    page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):font})})
    stream=DecodedStreamObject();stream.set_data(('BT /F1 12 Tf 15 200 Td ('+line+') Tj ET').encode())
    page[NameObject('/Contents')]=w._add_object(stream)
w.write(str(out/'material.pdf'))
deck=Presentation()
for line in lines:
    slide=deck.slides.add_slide(deck.slide_layouts[6]);slide.shapes.add_textbox(Inches(.5),Inches(1),Inches(9),Inches(2)).text=line

deck.save(str(out/'material.pptx'))
