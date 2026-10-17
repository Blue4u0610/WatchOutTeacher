"""Run parsers in a bounded subprocess. stdout is JSON, never native messaging."""
import json
import sys
import zipfile
from pathlib import Path
from common import AppError

MAX_TEXT=3_000_000

def parse(path, suffix):
    units=[]; warnings=[]
    if suffix in ['.txt','.md']:
        try: text=Path(path).read_text(encoding='utf-8-sig')
        except UnicodeDecodeError: raise AppError('encoding','文本必须为 UTF-8 编码。')
        if '\0' in text: raise AppError('invalid_document','文件包含二进制数据。')
        units=[{'number':1,'kind':'text','text':text}]
    elif suffix=='.pdf':
        from pypdf import PdfReader
        try:
            reader=PdfReader(path,strict=False)
            if reader.is_encrypted: raise AppError('encrypted','PDF 已加密；请先在本机解密后导入。')
            if len(reader.pages)>1500: raise AppError('too_large','PDF 超过 1,500 页，请拆分后导入。')
            for n,page in enumerate(reader.pages,1):
                text=page.extract_text() or ''
                if not text.strip(): warnings.append(f'PDF 文件页 {n} 无可提取文字（空白或扫描页，需要 OCR）。')
                units.append({'number':n,'kind':'page','text':text})
        except AppError: raise
        except Exception: raise AppError('invalid_document','PDF 损坏或文字提取失败。')
        if not any(u['text'].strip() for u in units): raise AppError('ocr_required','PDF 没有可提取文字，可能是扫描件，需要先 OCR。')
    elif suffix=='.pptx':
        from pptx import Presentation
        try:
            with zipfile.ZipFile(path) as z:
                if sum(i.file_size for i in z.infolist())>150*1024*1024: raise AppError('too_large','PPTX 解压体积超过 150 MB。')
            deck=Presentation(path)
            if len(deck.slides)>1500: raise AppError('too_large','PPTX 超过 1,500 张幻灯片。')
            def shape_text(shapes):
                texts=[]
                for shape in shapes:
                    if shape.has_text_frame: texts.append(shape.text)
                    if shape.has_table: texts.extend('\t'.join(c.text for c in row.cells) for row in shape.table.rows)
                    if hasattr(shape,'shapes'): texts.extend(shape_text(shape.shapes))
                return texts
            for n,slide in enumerate(deck.slides,1):
                text='\n'.join(shape_text(slide.shapes))
                if not text.strip(): warnings.append(f'幻灯片 {n} 无可提取文字；图形、图片与演讲者备注未导入。')
                units.append({'number':n,'kind':'slide','text':text})
            warnings.append('PPTX 仅解析可读文字与表格，不解析图片、图表语义、公式对象或演讲者备注。')
        except AppError: raise
        except Exception: raise AppError('invalid_document','PPTX 损坏、加密或结构不受支持。')
    else: raise AppError('unsupported','只支持 PDF、TXT、Markdown 和 PPTX；旧版 .ppt 请另存为 .pptx。')
    if not any(u['text'].strip() for u in units): raise AppError('empty','文件没有可索引的正文。')
    if sum(len(u['text']) for u in units)>MAX_TEXT: raise AppError('too_large','正文超过 300 万字符，请拆分资料；此次未导入。')
    return {'units':units,'warnings':warnings[:30]+([f'另有 {len(warnings)-30} 项页面警告。'] if len(warnings)>30 else [])}

if __name__=='__main__':
    try:
        # Windows limits are applied by the parent's Job Object before this process runs.
        if sys.platform != 'win32':
            import resource
            resource.setrlimit(resource.RLIMIT_CPU,(35,35))
            if sys.platform != 'darwin':
                resource.setrlimit(resource.RLIMIT_AS,(1536*1024*1024,1536*1024*1024))
        result=parse(sys.argv[1],sys.argv[2])
        print(json.dumps({'ok':True,'data':result},ensure_ascii=False))
    except AppError as e: print(json.dumps({'ok':False,'error':{'code':e.code,'message':e.message}}))
    except Exception: print(json.dumps({'ok':False,'error':{'code':'parse_failed','message':'解析失败；文件可能损坏或超出资源限制。'}}))
